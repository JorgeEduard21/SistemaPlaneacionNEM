"""Conversión del Markdown de la planeación a un documento Word (.docx)."""
import io
import re
from typing import List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from models import Planeacion

_COLOR_TITULO = RGBColor(0xC2, 0x5A, 0x38)
_COLOR_SECCION = RGBColor(0x49, 0x75, 0x5B)
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|__[^_]+__|_[^_]+_)")


def _agregar_texto(parrafo, texto: str) -> None:
    """Agrega texto respetando **negritas** e *itálicas*."""
    for parte in _INLINE.split(texto):
        if not parte:
            continue
        if (parte.startswith("**") and parte.endswith("**")) or (parte.startswith("__") and parte.endswith("__")):
            parrafo.add_run(parte[2:-2]).bold = True
        elif (parte.startswith("*") and parte.endswith("*")) or (parte.startswith("_") and parte.endswith("_")):
            parrafo.add_run(parte[1:-1]).italic = True
        else:
            parrafo.add_run(parte)


def _sombrear(celda, color_hex: str) -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def _es_separador(fila: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$", fila))


def _celdas(fila: str) -> List[str]:
    fila = fila.strip()
    if fila.startswith("|"):
        fila = fila[1:]
    if fila.endswith("|"):
        fila = fila[:-1]
    return [c.strip() for c in fila.split("|")]


def _agregar_tabla(doc, filas: List[str]) -> None:
    datos = [_celdas(f) for f in filas if not _es_separador(f)]
    if not datos:
        return
    columnas = max(len(f) for f in datos)
    tabla = doc.add_table(rows=len(datos), cols=columnas)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(datos):
        for j in range(columnas):
            celda = tabla.cell(i, j)
            celda.text = ""
            p = celda.paragraphs[0]
            _agregar_texto(p, fila[j] if j < len(fila) else "")
            for run in p.runs:
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
            if i == 0:
                _sombrear(celda, "F3E6DC")
    doc.add_paragraph()


def markdown_a_docx(planeacion: Planeacion) -> bytes:
    doc = Document()

    for seccion in doc.sections:
        seccion.top_margin = seccion.bottom_margin = Cm(2)
        seccion.left_margin = seccion.right_margin = Cm(2.2)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    # Encabezado con datos del docente
    perfil = planeacion.perfil
    encabezado = " | ".join(x for x in [perfil.escuela, perfil.nombre_docente, perfil.ciclo_escolar] if x)
    if encabezado:
        p = doc.sections[0].header.paragraphs[0]
        p.text = encabezado
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    lineas = planeacion.contenido.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()
        sin_espacios = linea.strip()

        if not sin_espacios or re.match(r"^(-{3,}|\*{3,}|_{3,})$", sin_espacios):
            i += 1
            continue

        # Tablas
        if sin_espacios.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i])
                i += 1
            _agregar_tabla(doc, bloque)
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", sin_espacios)
        if m:
            nivel = len(m.group(1))
            texto = m.group(2).strip().strip("*")
            h = doc.add_heading(texto, level=min(nivel, 4) if nivel > 1 else 0)
            for run in h.runs:
                run.font.color.rgb = _COLOR_TITULO if nivel == 1 else _COLOR_SECCION
            i += 1
            continue

        m = re.match(r"^(\s*)[-*+]\s+(.*)$", linea)
        if m:
            sangria = len(m.group(1).replace("\t", "    ")) // 2
            estilo_lista = "List Bullet 2" if sangria >= 1 else "List Bullet"
            texto = m.group(2)
            casilla = re.match(r"^\[( |x|X)\]\s*(.*)$", texto)
            if casilla:
                texto = ("☑ " if casilla.group(1).lower() == "x" else "☐ ") + casilla.group(2)
            _agregar_texto(doc.add_paragraph(style=estilo_lista), texto)
            i += 1
            continue

        m = re.match(r"^(\s*)\d+[.)]\s+(.*)$", linea)
        if m:
            sangria = len(m.group(1).replace("\t", "    ")) // 2
            _agregar_texto(doc.add_paragraph(style="List Number 2" if sangria >= 1 else "List Number"), m.group(2))
            i += 1
            continue

        if sin_espacios.startswith(">"):
            p = doc.add_paragraph()
            _agregar_texto(p, sin_espacios.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        _agregar_texto(doc.add_paragraph(), sin_espacios)
        i += 1

    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


def nombre_archivo(planeacion: Planeacion) -> str:
    base = re.sub(r"[^\w\s-]", "", planeacion.titulo, flags=re.UNICODE).strip()
    base = re.sub(r"\s+", "_", base)[:60] or "planeacion"
    return f"{base}_{planeacion.creada[:10]}.docx"

